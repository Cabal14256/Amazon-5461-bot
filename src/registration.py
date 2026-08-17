#!/usr/bin/env python3
"""
5461 申请登记表模块

自动记录每次 5461 申请的详细信息
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, Optional


class RegistrationManager:
    """5461 申请登记管理器"""
    
    # 默认登记表路径（工作区内）
    DEFAULT_REGISTRY_PATH = Path(__file__).parent.parent / 'data' / '5461申请登记表.xlsx'
    
    # 表头定义
    HEADERS = [
        '序号',
        '申请日期',
        '站点',
        '账号',
        '品牌',
        'SKU',
        'Case ID',
        '申请状态',
        '备注',
    ]
    
    # 状态选项
    STATUS_OPTIONS = ['申请中', '已通过', '假过', '已拒绝', '待补充材料', '待人工复核', '已取消']
    
    def __init__(self, registry_path: Optional[str] = None):
        self.registry_path = Path(registry_path or self.DEFAULT_REGISTRY_PATH)
        self._ensure_registry_exists()
    
    def _ensure_registry_exists(self) -> None:
        """确保登记表存在，不存在则创建"""
        if self.registry_path.exists():
            return
        
        print(f"[Registry] 创建新的登记表: {self.registry_path}")
        
        # 创建新的工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "5461申请记录"
        
        # 设置表头
        for col, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # 设置列宽
        column_widths = {
            'A': 8,   # 序号
            'B': 15,  # 申请日期
            'C': 12,  # 站点
            'D': 20,  # 账号
            'E': 15,  # 品牌
            'F': 25,  # SKU
            'G': 18,  # Case ID
            'H': 12,  # 申请状态
            'I': 30,  # 备注
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # 冻结首行
        ws.freeze_panes = 'A2'
        
        # 保存
        self.registry_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.registry_path)
        print(f"[Registry] 登记表创建完成")
    
    def add_record(self, record: Dict[str, Any]) -> int:
        """
        添加一条申请记录
        
        Args:
            record: 记录字典，包含以下字段:
                - site: 站点 (如 US, UK, DE)
                - account: 账号 (如 us_store_541)
                - brand: 品牌名
                - sku: SKU
                - case_id: Case ID
                - status: 状态 (默认 '申请中')
                - notes: 备注 (可选)
                
        Returns:
            新增记录的行号
        """
        # 加载工作簿
        wb = openpyxl.load_workbook(self.registry_path)
        ws = wb.active
        
        # 找到最后一行
        last_row = ws.max_row
        new_row = last_row + 1
        
        # 获取当前时间
        now = datetime.now()
        
        # 填充数据
        ws.cell(row=new_row, column=1, value=new_row - 1)  # 序号
        ws.cell(row=new_row, column=2, value=now.strftime('%Y-%m-%d %H:%M'))  # 申请日期
        ws.cell(row=new_row, column=3, value=record.get('site', ''))  # 站点
        ws.cell(row=new_row, column=4, value=record.get('account', ''))  # 账号
        ws.cell(row=new_row, column=5, value=record.get('brand', ''))  # 品牌
        ws.cell(row=new_row, column=6, value=record.get('sku', ''))  # SKU
        ws.cell(row=new_row, column=7, value=record.get('case_id', ''))  # Case ID
        ws.cell(row=new_row, column=8, value=record.get('status', '申请中'))  # 申请状态
        ws.cell(row=new_row, column=9, value=record.get('notes', ''))  # 备注
        
        # 设置单元格样式
        for col in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=new_row, column=col)
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
            # 状态列根据状态设置颜色
            if col == 8:  # 申请状态列
                status = record.get('status', '申请中')
                if status == '已通过':
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    cell.font = Font(color="006100")
                elif status == '假过':
                    cell.fill = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
                    cell.font = Font(color="9C5700")
                elif status == '已拒绝':
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    cell.font = Font(color="9C0006")
                elif status == '申请中':
                    cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                    cell.font = Font(color="9C5700")
                elif status in ('待补充材料', '待人工复核'):
                    cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                    cell.font = Font(color="9C5700")
        
        # 设置行高
        ws.row_dimensions[new_row].height = 20
        
        # 保存
        wb.save(self.registry_path)
        print(f"[Registry] 记录已添加 (行 {new_row}): {record.get('brand')} / {record.get('case_id', 'N/A')}")
        
        return new_row
    
    def update_status(self, case_id: str, new_status: str, notes: str = "") -> bool:
        """
        更新申请状态
        
        Args:
            case_id: Case ID
            new_status: 新状态
            notes: 备注（可选，会追加到现有备注）
            
        Returns:
            是否成功更新
        """
        if new_status not in self.STATUS_OPTIONS:
            print(f"[Registry] [警告] 未知状态: {new_status}，有效状态: {self.STATUS_OPTIONS}")
        
        wb = openpyxl.load_workbook(self.registry_path)
        ws = wb.active
        
        # 查找 Case ID
        for row in range(2, ws.max_row + 1):
            cell_value = ws.cell(row=row, column=7).value  # Case ID 列
            if cell_value and str(cell_value) == str(case_id):
                # 更新状态
                status_cell = ws.cell(row=row, column=8)
                status_cell.value = new_status
                
                # 更新颜色
                if new_status == '已通过':
                    status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    status_cell.font = Font(color="006100")
                elif new_status == '假过':
                    status_cell.fill = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
                    status_cell.font = Font(color="9C5700")
                elif new_status == '已拒绝':
                    status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    status_cell.font = Font(color="9C0006")
                elif new_status == '申请中':
                    status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                    status_cell.font = Font(color="9C5700")
                elif new_status in ('待补充材料', '待人工复核'):
                    status_cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                    status_cell.font = Font(color="9C5700")
                
                # 更新备注
                if notes:
                    notes_cell = ws.cell(row=row, column=9)
                    existing_notes = notes_cell.value or ""
                    if existing_notes:
                        notes_cell.value = f"{existing_notes}\n{datetime.now().strftime('%Y-%m-%d')}: {notes}"
                    else:
                        notes_cell.value = f"{datetime.now().strftime('%Y-%m-%d')}: {notes}"
                
                wb.save(self.registry_path)
                print(f"[Registry] 状态已更新: Case {case_id} -> {new_status}")
                return True
        
        print(f"[Registry] [警告] 未找到 Case ID: {case_id}")
        return False
    
    def find_record(self, case_id: Optional[str] = None, brand: Optional[str] = None, 
                    account: Optional[str] = None) -> list:
        """
        查找记录
        
        Args:
            case_id: Case ID（可选）
            brand: 品牌名（可选）
            account: 账号（可选）
            
        Returns:
            匹配的记录列表
        """
        wb = openpyxl.load_workbook(self.registry_path, data_only=True)
        ws = wb.active
        
        results = []
        for row in range(2, ws.max_row + 1):
            record = {
                'row': row,
                'seq': ws.cell(row=row, column=1).value,
                'date': ws.cell(row=row, column=2).value,
                'site': ws.cell(row=row, column=3).value,
                'account': ws.cell(row=row, column=4).value,
                'brand': ws.cell(row=row, column=5).value,
                'sku': ws.cell(row=row, column=6).value,
                'case_id': ws.cell(row=row, column=7).value,
                'status': ws.cell(row=row, column=8).value,
                'notes': ws.cell(row=row, column=9).value,
            }
            
            # 匹配条件
            match = True
            if case_id and str(record['case_id']) != str(case_id):
                match = False
            if brand and str(record.get('brand', '')).lower() != brand.lower():
                match = False
            if account and str(record.get('account', '')).lower() != account.lower():
                match = False
            
            if match:
                results.append(record)
        
        return results
    
    def get_statistics(self) -> Dict[str, Any]:
        """获取统计信息"""
        wb = openpyxl.load_workbook(self.registry_path, data_only=True)
        ws = wb.active
        
        stats = {
            'total': 0,
            'by_status': {status: 0 for status in self.STATUS_OPTIONS},
            'by_site': {},
            'by_brand': {},
        }
        
        for row in range(2, ws.max_row + 1):
            status = ws.cell(row=row, column=8).value
            site = ws.cell(row=row, column=3).value
            brand = ws.cell(row=row, column=5).value
            
            if status:
                stats['total'] += 1
                if status in stats['by_status']:
                    stats['by_status'][status] += 1
            
            if site:
                stats['by_site'][site] = stats['by_site'].get(site, 0) + 1
            
            if brand:
                stats['by_brand'][brand] = stats['by_brand'].get(brand, 0) + 1
        
        return stats


if __name__ == "__main__":
    # 测试代码
    print("="*60)
    print("5461 申请登记表测试")
    print("="*60)
    
    reg = RegistrationManager()
    
    # 添加测试记录
    test_record = {
        'site': 'US',
        'account': 'us_store_541',
        'brand': 'TestBrand',
        'sku': '541-US-TestBrand-ABC123',
        'case_id': '19999999999',
        'status': '申请中',
        'notes': '自动提交测试',
    }
    
    row = reg.add_record(test_record)
    print(f"\n添加记录到行: {row}")
    
    # 查找记录
    results = reg.find_record(case_id='19999999999')
    print(f"\n查找结果: {len(results)} 条")
    for r in results:
        print(f"  - {r['brand']} / {r['case_id']} / {r['status']}")
    
    # 更新状态
    reg.update_status('19999999999', '已通过', '测试通过')
    
    # 获取统计
    stats = reg.get_statistics()
    print(f"\n统计信息:")
    print(f"  总计: {stats['total']}")
    print(f"  按状态: {stats['by_status']}")
